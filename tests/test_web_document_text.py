"""Word / PDF 简历正文抽取测试。

自造最小 .docx / .pdf 样本，覆盖抽取结果、失败提示与上传接口。
"""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient

from boss_web.api.resume import MAX_BINARY_BYTES
from boss_web.app import create_app
from boss_web.services import document_text as dt


def make_docx(*, with_table: bool = True, empty: bool = False) -> bytes:
    """造一份带段落和表格的 .docx 样本。"""
    from docx import Document

    document = Document()
    if not empty:
        document.add_paragraph("张三")
        document.add_paragraph("13800138000 · zhangsan@example.com")
        if with_table:
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "工作经历"
            table.cell(0, 1).text = "示例科技 · 后端工程师"
            table.cell(1, 0).text = "技能"
            table.cell(1, 1).text = "Python、FastAPI"
        document.add_paragraph("自我评价：热爱工程化。")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_pdf(text: str = "Zhang San", *, blank: bool = False) -> bytes:
    """手搓一份只含单行文字的最小 PDF 样本。"""
    content = b"" if blank else f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET".encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(content)).encode()
        + b" >>\nstream\n"
        + content
        + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    startxref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\n"
        f"startxref\n{startxref}\n%%EOF\n"
    ).encode()
    return bytes(out)


def make_encrypted_pdf(password: str = "s3cret") -> bytes:
    """造一份带密码的加密 PDF 样本。"""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt(password)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


@pytest.fixture
def client():
    """提供不抛服务端异常的测试客户端。"""
    return TestClient(create_app(), raise_server_exceptions=False)


@pytest.mark.parametrize(
    "name,expected",
    [
        ("a.docx", ".docx"),
        ("A.DOCX", ".docx"),
        ("我的简历.Pdf", ".pdf"),
        ("resume.md", ".md"),
        ("没有后缀", ""),
        (".hidden", ""),
        ("", ""),
    ],
)
def test_ext_of(name, expected):
    """验后缀提取统一转小写并正确处理无后缀。"""
    assert dt.ext_of(name) == expected


def test_allowed_ext_covers_word_and_pdf():
    """验允许后缀覆盖 Word 与 PDF。"""
    assert {".docx", ".pdf"} <= dt.ALLOWED_EXT


def test_extract_text_defers_markdown_and_text_to_caller():
    """验纯文本不走抽取、交回调用方。"""
    assert dt.extract_text("## 技能".encode(), "a.md") is None
    assert dt.extract_text("普通文本".encode("gbk"), "a.txt") is None
    assert dt.extract_text(b"xx", "a.exe") is None


def test_docx_extracts_paragraphs_and_tables():
    """验 docx 段落与表格内容都能抽出。"""
    text = dt.extract_text(make_docx(), "简历.docx")

    assert "张三" in text
    assert "13800138000" in text
    assert "示例科技" in text
    assert "Python、FastAPI" in text
    assert "热爱工程化" in text


def test_docx_keeps_document_order():
    """验 docx 段落与表格保持原文顺序。"""
    text = dt.extract_text(make_docx(), "简历.docx")
    assert text.index("张三") < text.index("工作经历") < text.index("热爱工程化")


def test_docx_empty_raises_friendly_error():
    """验空 docx 抛出可读的友好错误。"""
    with pytest.raises(dt.DocumentTextError) as ei:
        dt.extract_text(make_docx(empty=True), "空.docx")
    assert "没有文字" in str(ei.value)


def test_old_doc_format_is_told_to_convert():
    """验旧 .doc 格式提示用户另存为。"""
    with pytest.raises(dt.DocumentTextError) as ei:
        dt.extract_text(b"\xd0\xcf\x11\xe0old binary", "老简历.docx")
    assert "另存为" in str(ei.value)
    assert ".docx" in str(ei.value)


def test_corrupt_docx_does_not_leak_exception_class():
    """验损坏 docx 不泄漏异常类名。"""
    with pytest.raises(dt.DocumentTextError) as ei:
        dt.extract_text(b"not a zip at all", "坏.docx")
    message = str(ei.value)
    assert "PackageNotFoundError" not in message
    assert "BadZipFile" not in message


def test_pdf_extracts_text():
    """验 PDF 正文能抽出。"""
    text = dt.extract_text(make_pdf("Zhang San Python"), "简历.pdf")
    assert "Zhang San" in text


def test_scanned_pdf_gets_scan_specific_hint():
    """验扫描件给出扫描专属提示。"""
    with pytest.raises(dt.DocumentTextError) as ei:
        dt.extract_text(make_pdf(blank=True), "扫描件.pdf")
    message = str(ei.value)
    assert "扫描" in message
    assert ".txt" in message


def test_encrypted_pdf_tells_user_to_remove_password():
    """验加密 PDF 提示去除密码。"""
    with pytest.raises(dt.DocumentTextError) as ei:
        dt.extract_text(make_encrypted_pdf(), "加密.pdf")
    assert "密码" in str(ei.value)


def test_corrupt_pdf_does_not_leak_exception_class():
    """验损坏 PDF 不泄漏异常类名。"""
    with pytest.raises(dt.DocumentTextError) as ei:
        dt.extract_text(b"%PDF-1.4 garbage", "坏.pdf")
    assert "PdfReadError" not in str(ei.value)


def test_tidy_collapses_blank_runs():
    """验 `_tidy` 压缩连续空行与行尾空白。"""
    assert dt._tidy("a\n\n\n\n\nb") == "a\n\nb"
    assert dt._tidy("a  \nb") == "a\nb"


def test_upload_docx(client):
    """验上传 docx 抽出正文并原样落库。"""
    r = client.post(
        "/api/resume/upload",
        files={
            "file": (
                "我的简历.docx",
                make_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert r.status_code == 200
    body = r.json()
    assert "张三" in body["raw"], "抽出来的正文要原样落库，解析接口读的就是它"
    assert "示例科技" in body["raw"]
    assert body["resume_id"].startswith("rs_")
    client.delete(f"/api/resume/item/{body['resume_id']}")


def test_upload_pdf(client):
    """验上传 PDF 成功返回简历编号。"""
    r = client.post(
        "/api/resume/upload",
        files={"file": ("简历.pdf", make_pdf("Li Si"), "application/pdf")},
    )
    assert r.status_code == 200
    body = r.json()
    assert body["resume_id"].startswith("rs_")
    client.delete(f"/api/resume/item/{body['resume_id']}")


def test_upload_broken_docx_returns_actionable_422(client):
    """验损坏 docx 返回可操作的 422。"""
    r = client.post(
        "/api/resume/upload",
        files={"file": ("简历.docx", b"garbage", "application/octet-stream")},
    )
    assert r.status_code == 422
    assert "另存为" in r.json()["message"]


def test_reject_message_mentions_word_and_pdf(client):
    """验拒绝消息点名 Word 与 PDF。"""
    r = client.post(
        "/api/resume/upload",
        files={"file": ("resume.doc", b"xx", "application/msword")},
    )
    assert r.status_code == 422
    message = r.json()["message"]
    assert "Word" in message and "PDF" in message


def test_binary_resumes_get_a_larger_size_ceiling(client, monkeypatch):
    """验二进制简历有更大的体积上限。"""
    assert MAX_BINARY_BYTES > 2 * 1024 * 1024

    r = client.post(
        "/api/resume/upload",
        files={"file": ("大简历.docx", b"x" * (3 * 1024 * 1024), "application/octet-stream")},
    )
    assert r.status_code == 422
    assert "超过" not in r.json()["message"]


def test_text_resume_still_capped_at_2mb(client):
    """验纯文本简历仍受 2MB 上限约束。"""
    r = client.post(
        "/api/resume/upload",
        files={"file": ("big.txt", b"x" * (3 * 1024 * 1024), "text/plain")},
    )
    assert r.status_code == 422
    assert "超过" in r.json()["message"]
